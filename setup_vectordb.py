import os

DOCS_FOLDER = "./superman_docs"
VECTOR_DB_FOLDER = "./superman_vectordb"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

def build_vectordb():
    from pptx import Presentation
    from lxml import etree
    from langchain_core.documents import Document
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    from langchain_community.embeddings import HuggingFaceEmbeddings
    from langchain_community.vectorstores import Chroma
    from langchain_community.document_loaders import PyPDFLoader, Docx2txtLoader

    docs = []
    for file in os.listdir(DOCS_FOLDER):
        path = os.path.join(DOCS_FOLDER, file)
        try:
            if file.endswith(".pptx"):
                prs = Presentation(path)
                for i, slide in enumerate(prs.slides):
                    xml = slide._element.xml
                    tree = etree.fromstring(xml.encode())
                    texts = tree.findall(".//a:t", NS)
                    text_parts = [
                        t.text.strip()
                        for t in texts
                        if t.text and t.text.strip()
                    ]
                    full_text = "\n".join(text_parts).strip()
                    if full_text:
                        docs.append(Document(
                            page_content=full_text,
                            metadata={"source": file, "slide": i + 1}
                        ))
            elif file.endswith(".pdf"):
                loader = PyPDFLoader(path)
                docs.extend(loader.load())
            elif file.endswith(".docx"):
                loader = Docx2txtLoader(path)
                docs.extend(loader.load())
        except Exception as e:
            print(f"Error loading {file}: {e}")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=150
    )
    chunks = splitter.split_documents(docs)

    embeddings = HuggingFaceEmbeddings(
        model_name="sentence-transformers/all-MiniLM-L6-v2"
    )

    Chroma.from_documents(
        documents=chunks,
        embedding=embeddings,
        persist_directory=VECTOR_DB_FOLDER
    )
    print(f"VectorDB built with {len(chunks)} chunks")

if __name__ == "__main__":
    build_vectordb()